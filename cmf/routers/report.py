import os
import json
import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
import openpyxl

from DB.database import get_db
from DB.models.quality import MasterBoc, StageInspection
from DB.models.oms import Part, Order, Product

router = APIRouter(prefix="/reports", tags=["reports"])

_TEMPLATE_DIR = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
TEMPLATE_PATH = os.path.join(_TEMPLATE_DIR, "CMF- Inspection Rport(QMS).xlsx")

FIRST_BOC_ROW = 15
LAST_BOC_ROW = 29


def _master_boc_id_from_stage_bbox(bbox: str) -> int:
    if not bbox or not bbox.strip():
        return None
    try:
        o = json.loads(bbox)
        mid = o.get("master_boc_id")
        if mid is None:
            return None
        return int(mid)
    except (TypeError, ValueError, json.JSONDecodeError):
        return None


def _fmt_num(x) -> str:
    if x is None:
        return ""
    try:
        v = float(x)
        if v == int(v):
            return str(int(v))
        return str(v).rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        return str(x)


def _nz(s) -> str:
    if s is None:
        return ""
    return str(s).strip()


@router.get("/inspection-report")
def get_inspection_report(
    part_number: str = Query(..., description="oms.parts.part_number"),
    sales_order_id: int = Query(...),
    op_no: int = Query(...),
    db: Session = Depends(get_db),
):
    """Fill `CMF- Inspection Rport(QMS).xlsx` with data; layout and styling stay in the template file."""
    if not os.path.isfile(TEMPLATE_PATH):
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Template file not found: {TEMPLATE_PATH}",
        )

    part = (
        db.query(Part)
        .options(joinedload(Part.assembly))
        .filter(Part.part_number == part_number)
        .first()
    )
    if not part:
        raise HTTPException(status_code=404, detail="Part not found")

    order = db.query(Order).filter(Order.id == sales_order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    product = db.query(Product).filter(Product.id == order.product_id).first()

    masters = (
        db.query(MasterBoc)
        .filter(
            MasterBoc.part_id == part_number,
            MasterBoc.sales_order_id == sales_order_id,
            MasterBoc.op_no == op_no,
        )
        .order_by(MasterBoc.id.asc())
        .all()
    )

    stage_rows = (
        db.query(StageInspection)
        .filter(
            StageInspection.part_id == part.id,
            StageInspection.sale_order_id == sales_order_id,
            StageInspection.op_no == op_no,
        )
        .all()
    )

    stage_by_master_id = {}
    for row in stage_rows:
        mid = _master_boc_id_from_stage_bbox(row.bbox)
        if mid is not None:
            stage_by_master_id.setdefault(mid, []).append(row)

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    sheet = wb.active

    # Clear previous run (value cells only)
    for r in range(FIRST_BOC_ROW, LAST_BOC_ROW + 1):
        for c in (1, 2, 5, 7, 9, 11, 13):
            sheet.cell(row=r, column=c).value = None
    for addr in ("D7", "J7", "O7", "D9", "J9", "O9", "D11", "J11", "O11"):
        sheet[addr].value = None

    now = datetime.datetime.now()
    date_disp = f"{now.month}/{now.day}/{now.year}"
    report_no = f"RPT-{sales_order_id}-{op_no}"
    assembly_disp = (
        part.assembly.assembly_name if part.assembly is not None else "Main"
    )

    sheet["D7"] = report_no
    sheet["J7"] = part.part_name
    sheet["O7"] = date_disp
    sheet["D9"] = str(order.sale_order_number)
    sheet["J9"] = part.part_number
    sheet["O9"] = "1 of 1"
    sheet["D11"] = product.product_name if product else ""
    sheet["J11"] = order.quantity if order.quantity is not None else ""
    sheet["O11"] = assembly_disp

    max_rows = LAST_BOC_ROW - FIRST_BOC_ROW + 1
    for i, m in enumerate(masters[:max_rows]):
        r = FIRST_BOC_ROW + i
        sheet.cell(row=r, column=1).value = i + 1
        sheet.cell(row=r, column=2).value = (
            f"{m.nominal} ({_fmt_num(m.uppertol)}/{_fmt_num(m.lowertol)})"
        )
        sheet.cell(row=r, column=5).value = m.zone

        g_val, i_val, k_val = "", "", ""
        items = stage_by_master_id.get(m.id, [])
        if items:
            st = items[0]
            meas = st.measurements or []
            m1 = _nz(meas[0]) if len(meas) > 0 else ""
            m2 = _nz(meas[1]) if len(meas) > 1 else ""
            m3 = _nz(meas[2]) if len(meas) > 2 else ""
            mean_s = _nz(st.measured_mean)
            if m1 or m2 or m3:
                g_val, i_val, k_val = m1, m2, m3
            elif mean_s:
                g_val = mean_s

        sheet.cell(row=r, column=7).value = g_val
        sheet.cell(row=r, column=9).value = i_val
        sheet.cell(row=r, column=11).value = k_val

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Inspection_Report_{part_number}_OP{op_no}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
